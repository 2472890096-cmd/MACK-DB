"""Frozen-input, same-sample edge comparison. No inference executable is called."""

# v0.1: repository-local IO configuration; scientific functions below are unchanged.
import sys as _path_sys
from pathlib import Path as _RepoPath
_repo_root = next(p for p in _RepoPath(__file__).resolve().parents if (p / "lib/mackdb_paths.py").is_file())
_path_sys.path.insert(0, str(_repo_root / "lib"))
from mackdb_paths import resolve_path
from pathlib import Path
from collections import Counter,defaultdict
import csv,datetime,hashlib,json,math,os,re,sys,traceback
import numpy as np
import openpyxl
from scipy.stats import spearmanr
ROOT=Path(__file__).resolve().parents[1]
STAGE=ROOT.parent; BASE=STAGE.parent; EXT=resolve_path('literature'); PREP=STAGE/'02_prepare_fastspar_input'; FSROOT=STAGE/'03_fastspar_result'
SPECROOT=Path(str(resolve_path('species_input')))
GUOWB=EXT/'raw/guo2023/1-s2.0-S193131282300029X-mmc8.xlsx'
WALLENWB=STAGE/'00_metadata_cache/Wallen_supplementary_datasets.xlsx'
WMAP=BASE/'figure56/PMID_36376318_taxonomy_rescue_overlap_v44/02_taxon_mapping_audit.tsv'
GMAP=EXT/'stage2b/03_external_taxa_to_gtdb_r226_stage2b.tsv'
NETWORKS=[('Guo_Control',91,377,'PRJNA751448'),('Guo_MECFS',106,377,'PRJNA751448'),('Guo_Common',197,377,'PRJNA751448'),('Wallen_PD',490,598,'PRJNA834801'),('Wallen_NHC',234,598,'PRJNA834801')]
QC=[]; INPUTS={}; DATA={}
NA='NA'
def now():return datetime.datetime.now(datetime.timezone.utc).isoformat()
def sha(p):
 with Path(p).open('rb') as f:return hashlib.file_digest(f,'sha256').hexdigest()
def register(p,role='frozen_input'):
 p=Path(p).resolve()
 if ROOT not in p.parents and str(p) not in INPUTS:INPUTS[str(p)]={'path':str(p),'role':role,'bytes':p.stat().st_size,'sha256_before':sha(p)}
 return p
def read(p):
 p=register(p)
 with p.open(newline='') as f:return list(csv.DictReader(f,delimiter='\t'))
def write(p,rows,fields=None):
 p=Path(p); assert ROOT in p.resolve().parents
 rows=list(rows);fields=fields or list(rows[0]); tmp=p.with_name(p.name+'.tmp')
 with tmp.open('w',newline='') as f:
  w=csv.DictWriter(f,fieldnames=fields,delimiter='\t',lineterminator='\n');w.writeheader();w.writerows(rows)
 os.replace(tmp,p)
def saveqc():write(ROOT/'qc/stage5B2_qc_summary.tsv',QC,['network','check','status','details','timestamp'])
def check(n,test,ok,detail='',failure='FAIL'):
 QC.append(dict(network=n,check=test,status='PASS' if ok else failure,details=str(detail),timestamp=now()));saveqc()
 if not ok and failure=='FAIL':raise RuntimeError(n+': '+test+': '+str(detail))
def info(n,test,detail,warn=False):
 QC.append(dict(network=n,check=test,status='WARN' if warn else 'INFO',details=str(detail),timestamp=now()));saveqc()
def sign(v):return 'positive' if v>0 else 'negative' if v<0 else 'zero'
def pair(a,b):return tuple(sorted((a,b)))
def safe_div(a,b):return a/b if b else NA
def finite(v):return isinstance(v,(int,float,np.floating)) and math.isfinite(v)
def load_vocabulary():
 vs=read(ROOT/'00_inputs/gtdb_r226_species_vocabulary.tsv')
 enc={r['production_encoded_label']:r['gtdb_r226_species'] for r in vs};names={r['gtdb_r226_species'] for r in vs}
 check('ALL','R226_serialization_unique_inverse',len(enc)==len(names)==len(vs),len(vs))
 refs=json.loads((ROOT/'00_inputs/reference_manifest.json').read_text())
 for x in refs:INPUTS[x['path']]={'path':x['path'],'role':'frozen_R226_vocabulary_source','bytes':x['bytes'],'sha256_before':x['sha256']}
 return enc,names

def mapping_tables(names):
 guo=read(GMAP); wallen=read(WMAP)
 # Authoritative, study-specific frozen decisions have precedence, including unresolved decisions.
 gm={r['raw_taxon']:dict(target=r['gtdb_r226_species'] if r['mapping_status'].startswith('mapped_') else '',status=r['mapping_status'],source=str(GMAP),record=r['raw_taxon_id']) for r in guo if r['paper_id']=='Guo2023'}
 wm={r['external_taxon']:dict(target=r['gtdb_species'] if r['gtdb_status'].startswith('accepted') else '',status=r['gtdb_status'],source=str(WMAP),record=r['external_taxon']) for r in wallen}
 # Complete Wallen tested universe includes nodes absent from final edges. Reuse only existing
 # exact-name decisions for these extra nodes; do not introduce new taxonomic assignments.
 pool=defaultdict(list)
 for r in guo:
  key=' '.join(r['raw_taxon'].replace('_',' ').split()).casefold()
  if r['mapping_status'].startswith('mapped_'):
   pool[key].append((r['gtdb_r226_species'],str(GMAP),r['raw_taxon_id']))
 def resolve(study,raw):
  if study=='Guo':
   if raw in gm:return gm[raw]
   return dict(target='',status='no_frozen_decision',source=str(GMAP),record='')
  if raw in wm:return wm[raw]
  candidates=pool.get(' '.join(raw.replace('_',' ').split()).casefold(),[]);targets={r[0] for r in candidates}
  if len(targets)==1:return dict(target=next(iter(targets)),status='reused_existing_exact_name_decision_for_tested_universe',source=str(GMAP),record=';'.join(sorted({r[2] for r in candidates})))
  return dict(target='',status='no_frozen_decision' if not targets else 'conflicting_frozen_decisions',source=str(GMAP),record='')
 check('ALL','frozen_mapping_targets_in_R226',all(x['target'] in names for x in list(gm.values())+list(wm.values()) if x['target']))
 register(EXT/'mapping/03_external_taxa_to_gtdb_r226.tsv','taxonomy_pre_rescue_provenance')
 register(EXT/'mapping/04_all_external_edges_gtdb_r226.tsv','provenance_only_not_Guo_edge_input')
 register(EXT/'mapping/04_edge_mapping_audit.tsv','provenance_only')
 register(EXT/'scripts/02_harmonize_external_taxonomy.py','frozen_taxonomy_rules')
 return resolve

def published_inputs():
 pub={};universes={};evidence=[]
 w=openpyxl.load_workbook(register(GUOWB),read_only=True,data_only=True);rows=list(w['Table S7'].values)
 schema=[]
 for n,j,count,nodes in [('Guo_Common',1,297,100),('Guo_Control',4,210,99),('Guo_MECFS',7,223,100)]:
  check(n,'Guo_raw_headers',rows[25][j:j+3]==('Taxa1','Taxa2','Weight'),rows[9][j])
  records=[]
  for rownum,r in enumerate(rows[26:],27):
   vals=r[j:j+3]
   if all(v is None for v in vals):continue
   check(n,'raw_complete_row_'+str(rownum),all(v is not None for v in vals),failure='FAIL') if any(v is None for v in vals) else None
   a,b,v=vals;records.append(dict(network=n,raw_edge_id=n+':'+str(rownum),source_file=str(GUOWB),source_sheet='Table S7',source_row=rownum,source_network_label=rows[9][j],species1=a,species2=b,published_weight_or_r=float(v),published_r=NA,published_p=NA,published_sign=sign(float(v)),selection_rule='author final single-fit edge list; FlashWeave-S adjusted P < 0.05; no added weight filter'))
  check(n,'published_raw_edge_count',len(records)==count,len(records))
  check(n,'Guo_node_count_in_workbook',len({r[k] for r in records for k in ['species1','species2']})==nodes,nodes)
  check(n,'Guo_published_sign_counts',sum(r['published_sign']=='positive' for r in records)==rows[12][j] and sum(r['published_sign']=='negative' for r in records)==rows[13][j])
  pub[n]=records;write(ROOT/'01_published_networks'/f'{n}_published_raw.tsv',records)
  schema.append(dict(network=n,source_file=str(GUOWB),sheet='Table S7',network_label_row=10,header_row=26,source_column=j+1,target_column=j+2,weight_column=j+3,sign_column='derived from signed weight',pvalue_column='not supplied',selection_fields='Methods adjusted P<0.05; workbook final edge list',raw_edges=count))
 w.close()
 text=register(STAGE/'01_sample_rescue/cache/Guo_article_extracted.txt','published_methods').read_text()
 start=text.index('Microbiome co-abundance network analysis');end=text.index('Next, we use Flashweave',start)+500
 (ROOT/'00_inputs/Guo_analysis_universe_methods_excerpt.txt').write_text(text[start:end]+'\n')
 check('Guo_ALL','full_analysis_universe_explicit_author_evidence','100 species after filtration' in text and 'all 100 species (nodes) included in the analysis' in text)
 common_nodes={r[k] for r in pub['Guo_Common'] for k in ['species1','species2']}
 for n in pub:
  check(n,'same_100_species_tested_universe',len(common_nodes)==100 and all(r[k] in common_nodes for r in pub[n] for k in ['species1','species2']))
  universes[n]=sorted(common_nodes);evidence.append(dict(network=n,universe_type='complete_published_analysis_taxa',raw_taxa=100,evidence='Guo Methods explicitly states same 100 filtered species and Common includes all 100; Table S7 Common supplies identities; includes Control isolated node',source_file=str(GUOWB)))
 w=openpyxl.load_workbook(register(WALLENWB),read_only=True,data_only=True)
 derived=openpyxl.load_workbook(register(BASE/'figure56/Wallen_etal.xlsx'),read_only=True,data_only=True)
 dv=list(derived['derived_edge_list'].values);dr=[dict(zip(dv[0],r)) for r in dv[1:]];derived.close()
 for group,sheet,total in [('PD','Supplementary Data 5',697),('NHC','Supplementary Data 6',499)]:
  n='Wallen_'+group;records=[];nodes=set();sh=w[sheet];legend='';head=[];n_pfiltered=0
  for rownum,row in enumerate(sh.values,1):
   if rownum==1:legend=row[0];continue
   if rownum==3:head=list(row);continue
   if rownum<4 or not row[0]:continue
   a,b,mag,direction,p,assoc=row;nodes.update([a,b]);n_pfiltered+=1
   assert float(p)<0.05 and direction in ['+','-']
   if not abs(float(mag))>0.2:continue
   r=float(mag)*(1 if direction=='+' else -1)
   records.append(dict(network=n,raw_edge_id=n+':'+str(rownum),source_file=str(WALLENWB),source_sheet=sheet,source_row=rownum,source_network_label=group,species1=a,species2=b,published_weight_or_r=r,published_r=r,published_p=float(p),published_sign=sign(r),selection_rule='abs(r)>0.2 AND uncorrected P<0.05',PD_associated=assoc))
  check(n,'published_full_tested_universe_count',len(nodes)==total and f'All {total} species' in legend,len(nodes))
  expected=read(BASE/f'figure56/mackdb_overlap_validation_v3/PMID_36376318/{group}/00_published_unique_pairs.tsv')
  def key(a,b,r,p):return (pair(a,b),round(float(r),10),round(float(p),10))
  current=Counter(key(r['species1'],r['species2'],r['published_r'],r['published_p']) for r in records)
  old=Counter(key(r['species_a'],r['species_b'],r['correlation'],r['pvalue']) for r in expected)
  deriv=Counter(key(r['species_a'],r['species_b'],r['correlation'],r['pvalue']) for r in dr if r['cohort_or_group']==group)
  check(n,'published_raw_edge_count_and_component_provenance',current==old==deriv and len(records)==(1817 if group=='PD' else 2509),f'{len(records)}; exact raw workbook r/P/edge identities equal subgroup component and derived workbook; no union/source exclusion')
  pub[n]=records;universes[n]=sorted(nodes)
  write(ROOT/'01_published_networks'/f'{n}_published.tsv',records)
  evidence.append(dict(network=n,universe_type='complete_published_analysis_taxa',raw_taxa=total,evidence=f'All {total} tested species explicitly stated in legend; ALL P-filtered table endpoints (before |r| filtering) recover exactly {total} distinct names',source_file=str(WALLENWB)))
  schema.append(dict(network=n,source_file=str(WALLENWB),sheet=sheet,network_label_row=1,header_row=3,source_column=1,target_column=2,weight_column=3,sign_column=4,pvalue_column=5,selection_fields='P<0.05 table; then abs(r)>0.2',raw_edges=len(records)))
  (ROOT/'00_inputs'/f'{n}_published_universe_legend.txt').write_text(legend+'\nHeaders: '+json.dumps(head)+'\n')
 w.close()
 write(ROOT/'00_inputs/published_workbook_schema.tsv',schema)
 write(ROOT/'00_inputs/published_universe_evidence.tsv',evidence)
 return pub,universes

def load_fastspar(n,samples,taxa,project,encoding):
 rowmap=read(PREP/'qc'/f'{n}_taxa_row_map.tsv');samplemap=read(PREP/'qc'/f'{n}_sample_column_map.tsv')
 p=register(SPECROOT/'species'/f'{project}_species.matrix','original_abundance_species_dictionary')
 with p.open() as f: species=[line.rstrip('\n').split('\t') for line in f]
 check(n,'species_dictionary_ID_uniqueness',len(species)==taxa and len({r[0] for r in species})==taxa and all(len(r)==2 for r in species))
 sd=dict(species);out=[]
 ids=[r['count_species_id'] for r in rowmap]
 check(n,'FastSpar_taxa_row_map_integrity',len(rowmap)==taxa and len(set(ids))==taxa and [r['row_position_1based'] for r in rowmap]==list(map(str,range(1,taxa+1))))
 p=register(PREP/'count_matrix'/f'{n}_count.matrix')
 with p.open() as f:counts=list(csv.reader(f,delimiter='\t'))
 p=register(PREP/'abundance_subset'/f'{n}_abundance.matrix')
 with p.open() as f:abundance=list(csv.reader(f,delimiter='\t'))
 p=register(SPECROOT/'abundance'/f'{project}_abundance.matrix')
 with p.open() as f:original_ids=[r[0] for r in list(csv.reader(f,delimiter='\t'))[1:]]
 check(n,'count_abundance_original_ID_chain',ids==[r[0] for r in counts[1:]] and [r['abundance_species_id'] for r in rowmap]==[r[0] for r in abundance[1:]] and [r['original_species_id'] for r in rowmap]==original_ids==[r[0] for r in species])
 check(n,'sample_identity_chain',len(samplemap)==samples==len(counts[0])-1 and [r['count_sample_label'] for r in samplemap]==counts[0][1:] and [r['abundance_sample_label'] for r in samplemap]==abundance[0][1:] and len({r['run_accession'] for r in samplemap})==samples)
 for r in rowmap:
  raw=sd[r['original_species_id']];m=re.fullmatch(r'([0-9]+)_(.+)',raw)
  check(n,'production_taxon_label_'+r['count_species_id'],bool(m) and m[2] in encoding,raw) if not m or m[2] not in encoding else None
  out.append(dict(r,project=project,original_production_species_label=raw,production_species_numeric_prefix=m[1],gtdb_r226_species=encoding[m[2]],mapping_rule='exact inverse of numeric-prefix + R226 species space-to-underscore serialization; preserve GTDB suffixes',species_dictionary=str(p.parent.parent/'species'/f'{project}_species.matrix')))
 names=[r['gtdb_r226_species'] for r in out]
 check(n,'restored_species_unique',len(set(names))==taxa)
 matrices=[]
 for kind in ['correlation','pvalues']:
  p=register(FSROOT/n/f'{n}_{kind}.tsv')
  with p.open() as f:rr=list(csv.reader(f,delimiter='\t'))
  arr=np.asarray([r[1:] for r in rr[1:]],dtype=float)
  check(n,kind+'_IDs_order_dim_finite_symmetric',rr[0][1:]==ids==[r[0] for r in rr[1:]] and arr.shape==(taxa,taxa) and np.isfinite(arr).all() and np.array_equal(arr,arr.T))
  matrices.append(arr)
 check(n,'correlation_and_P_ranges',np.all(np.abs(matrices[0])<=1) and np.all((matrices[1]>=0)&(matrices[1]<=1)))
 register(FSROOT/n/f'{n}_covariance.tsv','frozen_result_not_used_for_edges')
 write(ROOT/'00_inputs'/f'{n}_restored_taxa.tsv',out)
 return names,matrices[0],matrices[1]

def harmonize(n,records,universe,resolve,fsnames):
 study=n.split('_')[0];taxa=sorted(set(universe)|{r[k] for r in records for k in ['species1','species2']});tm={t:resolve(study,t) for t in taxa}
 write(ROOT/'01_published_networks'/f'{n}_taxon_mapping_decisions.tsv',[dict(network=n,raw_taxon=t,gtdb_r226_species=tm[t]['target'] or NA,mapping_status=tm[t]['status'],source_file=tm[t]['source'],source_record=tm[t]['record'],in_published_final_edges=any(t in (r['species1'],r['species2']) for r in records)) for t in taxa])
 audited=[];groups=defaultdict(list)
 for r in records:
  a,b=tm[r['species1']],tm[r['species2']];q=dict(r,gtdb_taxon1=a['target'] or NA,gtdb_taxon2=b['target'] or NA,taxon1_mapping_status=a['status'],taxon2_mapping_status=b['status'])
  if not a['target'] or not b['target']:q['mapping_status']='taxonomy_unresolved'
  elif a['target']==b['target']:q['mapping_status']='self_loop_after_mapping'
  else:
   x,y=pair(a['target'],b['target']);q.update(taxon1=x,taxon2=y,mapping_status='mapped');groups[(x,y)].append(q)
  audited.append(q)
 edges=[]
 for (a,b),members in sorted(groups.items()):
  signs={r['published_sign'] for r in members};weights={r['published_weight_or_r'] for r in members};rvals={r['published_r'] for r in members};ps={r['published_p'] for r in members}
  edges.append(dict(network=n,taxon1=a,taxon2=b,published_sign=next(iter(signs)) if len(signs)==1 else NA,sign_status='consistent' if len(signs)==1 else 'conflict',published_weight_or_r=next(iter(weights)) if len(weights)==1 else NA,published_r=next(iter(rvals)) if len(rvals)==1 else NA,published_p=next(iter(ps)) if len(ps)==1 else NA,raw_record_count=len(members),raw_edge_ids=json.dumps([r['raw_edge_id'] for r in members]),published_weights_or_r=json.dumps(sorted(weights)),collapse_rule='pair presence retained; conflicting signs or nonunique coefficients stored as NA, never averaged'))
  for j,r in enumerate(members):r['mapping_status']='mapped' if j==0 else 'duplicate_undirected_pair_collapsed'
 fields=list(dict.fromkeys(k for r in audited for k in r));audited=[{k:r.get(k,NA) for k in fields} for r in audited]
 write(ROOT/'01_published_networks'/f'{n}_mapping_audit.tsv',audited)
 write(ROOT/'01_published_networks'/f'{n}_gtdb_r226_edges.tsv',edges)
 statuses=Counter(r['mapping_status'] for r in audited);rawends={r[k] for r in records for k in ['species1','species2']}
 stats=dict(network=n,raw_edges=len(records),unique_raw_endpoints=len(rawends),mapped_endpoints=sum(bool(tm[t]['target']) for t in rawends),unresolved_endpoints=sum(not tm[t]['target'] for t in rawends),mapped_endpoint_occurrences=sum(bool(tm[r[k]]['target']) for r in records for k in ['species1','species2']),unresolved_edge_records=statuses['taxonomy_unresolved'],self_loops_after_mapping=statuses['self_loop_after_mapping'],duplicate_undirected_pairs_collapsed=statuses['duplicate_undirected_pair_collapsed'],final_unique_gtdb_pairs=len(edges))
 write(ROOT/'01_published_networks'/f'{n}_mapping_summary.tsv',[stats])
 check(n,'taxonomy_edge_fate_balance',len(records)==stats['unresolved_edge_records']+stats['self_loops_after_mapping']+stats['duplicate_undirected_pairs_collapsed']+len(edges),stats)
 info(n,'taxonomy_mapping_completeness',f"{stats['mapped_endpoints']}/{len(rawends)} final-edge raw endpoints mapped; unresolved excluded before biological comparison",warn=stats['unresolved_endpoints']>0)
 check(n,'self_loop_removal_and_duplicate_collapse',all(e['taxon1']<e['taxon2'] for e in edges) and len({(e['taxon1'],e['taxon2']) for e in edges})==len(edges))
 mapped_universe={tm[t]['target'] for t in universe if tm[t]['target']};common=sorted(mapped_universe & set(fsnames))
 write(ROOT/'03_common_universe'/f'{n}_common_taxa.tsv',[dict(network=n,gtdb_r226_species=t,in_published_analysis_universe=True,in_fastspar_input=True) for t in common])
 write(ROOT/'03_common_universe'/f'{n}_published_analysis_taxa.tsv',[dict(network=n,raw_taxon=t,gtdb_r226_species=tm[t]['target'] or NA,mapping_status=tm[t]['status'],in_fastspar_input=tm[t]['target'] in fsnames) for t in universe])
 univ=dict(network=n,published_universe_type='complete_published_analysis_taxa',published_analysis_raw_taxa=len(universe),published_analysis_mapped_raw_taxa=sum(bool(tm[t]['target']) for t in universe),published_analysis_unresolved_raw_taxa=sum(not tm[t]['target'] for t in universe),published_analysis_unique_gtdb_taxa=len(mapped_universe),published_final_edge_endpoint_taxa=len(rawends),fastspar_input_taxa=len(fsnames),common_taxa=len(common),comparison_scope='intersection of mapped complete published analysis taxa and FastSpar input; not all taxa profiled by either pipeline',jaccard_type='common-mapped-analysis-universe Jaccard',full_published_tested_identity_evidence=str(ROOT/'00_inputs/published_universe_evidence.tsv'),unmapped_universe_handling='excluded; no new/fuzzy taxonomic assignments; frozen study-specific decisions take precedence')
 write(ROOT/'03_common_universe'/f'{n}_universe_audit.tsv',[univ])
 check(n,'common_universe_definition',set(common)==mapped_universe & set(fsnames),univ)
 info(n,'unresolved_analysis_universe_taxa',univ['published_analysis_unresolved_raw_taxa'],warn=univ['published_analysis_unresolved_raw_taxa']>0)
 return edges,audited,stats,common,univ

def passed(n,r,p,mode):
 if n.startswith('Wallen') and mode=='primary':return abs(r)>0.2,p<0.05
 return abs(r)>=0.3,p<=0.0001

def comparison(n,samples,names,corr,pvals,pub,audited,mapping,common,univ,mode='primary'):
 commonset=set(common);positions={t:i for i,t in enumerate(names)};i,j=np.triu_indices(len(names),1)
 rpass,ppass=passed(n,corr[i,j],pvals[i,j],mode); mask=rpass & ppass
 fs=[]
 for a,b in zip(i[mask],j[mask]):
  x,y=pair(names[a],names[b]);fs.append(dict(taxon1=x,taxon2=y,r=float(corr[a,b]),pvalue=float(pvals[a,b]),sign=sign(float(corr[a,b])),matrix_row_i=int(a),matrix_row_j=int(b),both_taxa_in_common_universe=x in commonset and y in commonset))
 fs.sort(key=lambda r:(r['taxon1'],r['taxon2']))
 label='primary' if n.startswith('Guo') else 'paper_matched' if mode=='primary' else 'mackdb_standard'
 fsfile=ROOT/'02_fastspar_networks'/f'{n}_fastspar_{label}.tsv';write(fsfile,fs,['taxon1','taxon2','r','pvalue','sign','matrix_row_i','matrix_row_j','both_taxa_in_common_universe'])
 eligible=[e for e in pub if e['taxon1'] in commonset and e['taxon2'] in commonset];restricted=[e for e in fs if e['both_taxa_in_common_universe']]
 pubfile=ROOT/'03_common_universe'/f'{n}_published_eligible_edges.tsv'
 if mode=='primary':write(pubfile,eligible,list(pub[0]))
 restrictedfile=ROOT/'03_common_universe'/f'{n}_fastspar_{label}_eligible_edges.tsv';write(restrictedfile,restricted,['taxon1','taxon2','r','pvalue','sign','matrix_row_i','matrix_row_j','both_taxa_in_common_universe'])
 ep={(e['taxon1'],e['taxon2']):e for e in eligible}; ef={(e['taxon1'],e['taxon2']):e for e in restricted};shared=sorted(ep.keys() & ef.keys());signs=Counter();sharedrows=[]
 for a,b in shared:
  p=ep[(a,b)];f=ef[(a,b)];ps=p['published_sign'];ss=f['sign'];conc=ps==ss if ps in ['positive','negative'] else NA;signs[(ps,ss)]+=1
  sharedrows.append(dict(network=n,analysis=mode,taxon1=a,taxon2=b,published_sign=ps,fastspar_sign=ss,sign_concordant=conc,published_weight_or_r=p['published_weight_or_r'],published_r=p['published_r'],fastspar_r=f['r'],fastspar_p=f['pvalue'],published_raw_records=p['raw_record_count']))
 sharedfile=ROOT/'05_sign_concordance'/f'{n}_{label}_shared_pairs.tsv';write(sharedfile,sharedrows,['network','analysis','taxon1','taxon2','published_sign','fastspar_sign','sign_concordant','published_weight_or_r','published_r','fastspar_r','fastspar_p','published_raw_records'])
 pp=signs[('positive','positive')];nn=signs[('negative','negative')];pn=signs[('positive','negative')];np_=signs[('negative','positive')]
 signstat=dict(network=n,analysis=mode,shared_edges=len(shared),positive_positive=pp,negative_negative=nn,positive_negative=pn,negative_positive=np_,same_sign=pp+nn,opposite_sign=pn+np_,unknown_or_conflicting_sign=len(shared)-pp-nn-pn-np_,sign_concordance=safe_div(pp+nn,len(shared)))
 fate=[]
 for e in audited:
  a,b=e.get('taxon1',NA),e.get('taxon2',NA);infs=a in positions and b in positions and a!=b;rv=pv=NA; rp=ppass_=False;sharededge=False;concordant=NA
  if e['mapping_status']=='taxonomy_unresolved':reason='taxonomy_unresolved'
  elif e['mapping_status']=='self_loop_after_mapping':reason='other'
  elif not infs:reason='taxon_not_in_fastspar_input'
  else:
   rv=float(corr[positions[a],positions[b]]);pv=float(pvals[positions[a],positions[b]]);rp,ppass_=passed(n,rv,pv,mode)
   sharededge=(a,b) in shared
   reason='' if sharededge else 'both_r_and_p_fail' if not rp and not ppass_ else 'abs_r_below_threshold' if not rp else 'p_above_threshold' if not ppass_ else 'other'
   if sharededge:concordant=e['published_sign']==sign(rv)
  fate.append(dict(network=n,analysis=mode,raw_edge_id=e['raw_edge_id'],source_row=e['source_row'],raw_species1=e['species1'],raw_species2=e['species2'],taxon1=a,taxon2=b,published_sign=e['published_sign'],published_weight_or_r=e['published_weight_or_r'],mapping_status=e['mapping_status'],both_taxa_in_fastspar=infs,fastspar_r=rv,fastspar_p=pv,fastspar_pass_threshold=bool(rp and ppass_),shared_edge=sharededge,sign_concordant=concordant,reason_if_not_recovered=reason,reason_detail='self_loop_after_mapping' if e['mapping_status']=='self_loop_after_mapping' else '',denominator_note='raw published edge record; duplicates count once in unique-pair metrics'))
 fatefile=ROOT/'04_overlap'/(f'{n}_edge_fate.tsv' if mode=='primary' else f'{n}_mackdb_standard_edge_fate.tsv');write(fatefile,fate)
 unique_fates=[]
 for e in pub:
  a,b=e['taxon1'],e['taxon2'];q=next(x for x in fate if x['taxon1']==a and x['taxon2']==b);unique_fates.append(dict(network=n,analysis=mode,taxon1=a,taxon2=b,eligible=a in commonset and b in commonset,raw_record_count=e['raw_record_count'],shared_edge=q['shared_edge'],reason_if_not_recovered=q['reason_if_not_recovered']))
 write(ROOT/'04_overlap'/f'{n}_{label}_unique_pair_fate.tsv',unique_fates)
 summary=dict(study=n.split('_')[0],network=n,comparison_type='same-sample_cross-framework' if n.startswith('Guo') else 'same-sample_cross-pipeline',samples=samples,published_method='FlashWeave-S' if n.startswith('Guo') else 'FastSpar',fastspar_method='MACK-DB MetaKSSD/GTDB R226 + FastSpar',analysis_label='Guo_primary_MACKDB_standard' if n.startswith('Guo') else 'Wallen_primary_paper_matched' if mode=='primary' else 'Wallen_sensitivity_MACKDB_standard',threshold='abs(r)>0.2 AND P<0.05' if n.startswith('Wallen') and mode=='primary' else 'abs(r)>=0.3 AND P<=0.0001',published_raw_edges=mapping['raw_edges'],published_final_gtdb_edges=len(pub),published_eligible_edges=len(ep),published_universe_type=univ['published_universe_type'],common_taxa=len(common),fastspar_edges_all_input_taxa=len(fs),fastspar_edges=len(ef),shared_edges=len(shared),published_edge_recovery=safe_div(len(shared),len(ep)),fastspar_to_published_overlap=safe_div(len(shared),len(ef)),jaccard=safe_div(len(shared),len(ep.keys()|ef.keys())),jaccard_type=univ['jaccard_type'],sign_concordance=signstat['sign_concordance'],published_eligible_edge_file=str(pubfile),fastspar_eligible_edge_file=str(restrictedfile),shared_pair_file=str(sharedfile),null_mean_shared=NA,null_enrichment=NA,null_z=NA,empirical_p=NA)
 check(n,mode+'_FastSpar_threshold_and_upper_triangle',len(fs)==int(mask.sum()) and all(x['matrix_row_i']<x['matrix_row_j'] and all(passed(n,x['r'],x['pvalue'],mode)) for x in fs))
 check(n,mode+'_pair_canonicalization',len({(r['taxon1'],r['taxon2']) for r in fs})==len(fs) and all(r['taxon1']<r['taxon2'] for r in fs))
 check(n,mode+'_shared_count_sign_and_fate',len(shared)==len([k for k in ep if k in ef]) and sum(q['shared_edge'] for q in unique_fates)==len(shared) and sum(signs.values())==len(shared))
 print('COMPARISON',json.dumps({k:summary[k] for k in ['network','analysis_label','published_raw_edges','published_final_gtdb_edges','published_eligible_edges','common_taxa','fastspar_edges','shared_edges','published_edge_recovery','jaccard','sign_concordance']}),flush=True)
 return summary,signstat

def main():
 (ROOT/'logs/state.json').write_text(json.dumps(dict(status='RUNNING',step='edge_comparison',timestamp=now()))+'\n')
 register(Path(__file__).resolve())
 encoding,names=load_vocabulary();resolve=mapping_tables(names);pub,universes=published_inputs()
 oldstatus=read(FSROOT/'qc/stage5B1B_run_status.tsv');check('ALL','Stage5B1B_COMPLETE_PASS',len(oldstatus)==5 and all(r['status']=='COMPLETE' and r['qc_status']=='PASS' for r in oldstatus))
 oldmanifest=read(FSROOT/'qc/stage5B1B_final_delivery_manifest.tsv');expected={r['path']:r['sha256'] for r in oldmanifest}
 primary=[];signstats=[];sensitivity=[];mapped_union=set()
 for n,s,t,project in NETWORKS:
  fsnames,corr,pvals=load_fastspar(n,s,t,project,encoding)
  for kind in ['correlation','covariance','pvalues']:
   p=FSROOT/n/f'{n}_{kind}.tsv';check(n,kind+'_frozen_delivery_hash',sha(p)==expected[str(p)])
  edges,audited,mapping,common,univ=harmonize(n,pub[n],universes[n],resolve,fsnames)
  result,ss=comparison(n,s,fsnames,corr,pvals,edges,audited,mapping,common,univ);primary.append(result);signstats.append(ss)
  if n.startswith('Wallen'):
   mapped_union.update((r['taxon1'],r['taxon2']) for r in edges)
   second,sss=comparison(n,s,fsnames,corr,pvals,edges,audited,mapping,common,univ,'sensitivity');sensitivity.extend([result.copy(),second]);signstats.append(sss)
   shared=read(result['shared_pair_file']);valid=[r for r in shared if r['published_r']!=NA and finite(float(r['published_r']))]
   rho=p=NA
   if len(valid)>=3 and len({r['published_r'] for r in valid})>1 and len({r['fastspar_r'] for r in valid})>1:
    r=spearmanr([float(x['published_r']) for x in valid],[float(x['fastspar_r']) for x in valid]);rho=float(r.statistic);p=float(r.pvalue)
   write(ROOT/'07_statistics'/f'{n}_shared_r_spearman.tsv',[dict(network=n,analysis='Wallen_primary_paper_matched',shared_edges=len(shared),pairs_with_unique_reliable_published_r=len(valid),excluded_nonunique_or_unavailable_r=len(shared)-len(valid),spearman_rho=rho,spearman_p=p,interpretation='same-sample cross-pipeline FastSpar; upstream abundance and inference parameters differ; shared-edge conditional association')])
 # Confirm only taxonomy assignments were reused from the historic merged reference, never merged edges.
 old=read(BASE/'figure56/PMID_36376318_taxonomy_rescue_overlap_v44/07_unique_gtdb_pairs.tsv')
 oldpairs={pair(r['gtdb_a'],r['gtdb_b']) for r in old}
 check('Wallen_ALL','frozen_taxonomy_pair_identity_reconciliation',mapped_union==oldpairs,f'component-derived mapped union={len(mapped_union)}, frozen reference={len(oldpairs)}; union used only for provenance check')
 write(ROOT/'07_statistics/stage5B2_primary_summary.tsv',primary)
 write(ROOT/'07_statistics/stage5B2_wallen_threshold_sensitivity.tsv',sensitivity)
 write(ROOT/'05_sign_concordance/stage5B2_sign_concordance_summary.tsv',signstats)
 write(ROOT/'00_inputs/input_manifest.tsv',INPUTS.values())
 (ROOT/'logs/state.json').write_text(json.dumps(dict(status='RUNNING',step='edge_comparison_complete_null_pending',timestamp=now()))+'\n')
 print('EDGE_COMPARISON_COMPLETE; NULL_PENDING',flush=True)

if __name__=='__main__':
 try:main()
 except Exception as e:
  (ROOT/'logs/failure.json').write_text(json.dumps(dict(timestamp=now(),error=repr(e),traceback=traceback.format_exc()),indent=2)+'\n')
  if INPUTS:write(ROOT/'00_inputs/input_manifest.tsv',INPUTS.values())
  raise
