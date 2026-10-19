#!/usr/bin/python3 -s
"""Phase 1 only: parse local published networks, preserve raw labels, and audit.

Run: python3 -s scripts/01_parse_external_networks.py
The fixed BASE is intentional. Nothing is written under raw. No network access,
taxonomy parsing/mapping, cross-network deduplication, or source exclusion occurs.
Only the Python standard library and the installed system openpyxl are needed.
Normal python3 invocation relaunches with -s to avoid this machine's incompatible
user-site NumPy/system-openpyxl combination; no packages or settings are changed.
"""

# v0.1: repository-local IO configuration; scientific functions below are unchanged.
import sys as _path_sys
from pathlib import Path as _RepoPath
_repo_root = next(p for p in _RepoPath(__file__).resolve().parents if (p / "lib/mackdb_paths.py").is_file())
_path_sys.path.insert(0, str(_repo_root / "lib"))
from mackdb_paths import resolve_path

import sys
import os
import site

if __name__ == '__main__' and site.ENABLE_USER_SITE:
    os.execv(sys.executable, [sys.executable, '-s', os.path.abspath(__file__)] + sys.argv[1:])

import csv
import hashlib
import json
import logging
import math
import tempfile
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

BASE = Path(str(resolve_path('literature')))
RAW = BASE / 'raw'
PARSED = BASE / 'parsed'
REPORTS = BASE / 'reports'
NA = 'NA'
FIELDS = ['paper_id', 'network_id', 'condition', 'source_file', 'taxon1_raw',
          'taxon2_raw', 'association_value', 'association_type', 'uncorrected_r',
          'corrected_r', 'pvalue', 'qvalue', 'sign', 'method', 'primary_status']
SHORT_FIELDS = [f for f in FIELDS if f not in ('uncorrected_r', 'corrected_r')]
GROUPS = ['MHNO', 'MHO', 'MUNO', 'MUO']
ZHOU_EXPECTED = {'MS': (6742, 786), 'HHC': (5688, 773)}
GUO_SPECS = [
    ('Common Network', 'Common Network (n=197)', 'Common', 'Common', 297, 'intermediate'),
    ('Control Network', 'Control Network (n=91)', 'Control', 'Control', 210, 'intermediate'),
    ('ME/CFS Network', 'ME/CFS Network (n=106)', 'MECFS', 'ME/CFS', 223, 'intermediate'),
    ('Control Unique Edge Network', 'Control Unique Edge Network', 'ControlUnique', 'Control', 91, 'validated_unique'),
    ('ME/CFS Unique Edge Network', 'ME/CFS Unique Edge Network', 'MECFSUnique', 'ME/CFS', 61, 'validated_unique'),
    ('Aggregated Control Network', 'Aggregated Control Network', 'AggControl', 'Control', 388, 'final_aggregated'),
    ('Aggregated ME/CFS Network', 'Aggregated ME/CFS Network', 'AggMECFS', 'ME/CFS', 358, 'final_aggregated'),
]
LACRUZ_EXPECTED = dict(zip(GROUPS, [690, 698, 514, 456]))
META_FIELDS = ['subject_id', 'run_accession', 'study_name', 'phenotype']
STATE = {'inventory': [], 'zhou': [], 'guo': [], 'lacruz': [], 'sheets': [],
         'warnings': [], 'outputs': [], 'status': 'RUNNING'}
LOG = logging.getLogger('phase1')


def sha256(path):
    digest = hashlib.sha256()
    with path.open('rb') as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b''):
            digest.update(block)
    return digest.hexdigest()


def write_tsv(path, rows, fields):
    assert RAW.resolve() not in path.resolve().parents, 'Writes to raw are prohibited'
    with path.open('w', encoding='utf-8', newline='') as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, delimiter='\t',
                                lineterminator='\n', extrasaction='ignore')
        writer.writeheader()
        for row in rows:
            writer.writerow({f: NA if row.get(f) is None else row.get(f, NA) for f in fields})
    # Exact read-back checks include every raw taxon string, weight and NA field.
    with path.open(encoding='utf-8', newline='') as handle:
        reader = csv.DictReader(handle, delimiter='\t')
        check(reader.fieldnames == fields, '{}: output header mismatch'.format(path.name))
        count = 0
        for count, observed in enumerate(reader, 1):
            check(count <= len(rows), '{}: excess output row'.format(path.name))
            expected = {f: NA if rows[count-1].get(f) is None else str(rows[count-1].get(f, NA)) for f in fields}
            check(observed == expected, '{}: read-back mismatch at row {}'.format(path.name, count+1))
        check(count == len(rows), '{}: output row-count mismatch'.format(path.name))


def qc_tsv(name, rows):
    fields = list(dict.fromkeys(key for row in rows for key in row))
    write_tsv(REPORTS / name, rows, fields)


def check(condition, message):
    # Actual assert, with optimized Python explicitly prohibited in main().
    assert condition, message


def num(value, context, probability=False, correlation=False):
    check(value is not None and value != '' and not isinstance(value, bool),
          '{}: missing/boolean numeric value'.format(context))
    result = float(value)
    check(math.isfinite(result), '{}: nonfinite value'.format(context))
    if probability:
        check(0 <= result <= 1, '{}: probability outside [0, 1]'.format(context))
    if correlation:
        check(-1 <= result <= 1, '{}: correlation outside [-1, 1]'.format(context))
    return result


def taxon(value, context):
    check(isinstance(value, str) and bool(value.strip()), '{}: missing/non-string taxon'.format(context))
    return value  # Never strip, split, rename, case-fold, or normalize a taxon.


def sign(value):
    return 'positive' if value > 0 else 'negative' if value < 0 else 'zero'


def edge(paper, network, condition, path, a, b, value, atype, method, status='primary', **extra):
    record = dict.fromkeys(FIELDS, NA)
    record.update(paper_id=paper, network_id=network, condition=condition,
                  source_file=path.name, taxon1_raw=taxon(a, network),
                  taxon2_raw=taxon(b, network), association_value=value,
                  association_type=atype, sign=sign(num(value, network)),
                  method=method, primary_status=status)
    record.update(extra)
    return record


def pair(row):
    # Sorting is an ephemeral QC key only. Output endpoint order is unchanged.
    return tuple(sorted((row['taxon1_raw'], row['taxon2_raw'])))


def edge_qc(rows):
    pairs = Counter(pair(row) for row in rows)
    signs = Counter(row['sign'] for row in rows)
    return dict(raw_edges=len(rows), unique_raw_taxa=len({row[f] for row in rows for f in ['taxon1_raw', 'taxon2_raw']}),
                positive_edges=signs['positive'], negative_edges=signs['negative'], zero_edges=signs['zero'],
                duplicate_raw_pairs=sum(n - 1 for n in pairs.values()),
                duplicate_pair_groups=sum(n > 1 for n in pairs.values()),
                self_loops=sum(row['taxon1_raw'] == row['taxon2_raw'] for row in rows))


def assert_simple(qc, network):
    check(qc['duplicate_raw_pairs'] == 0, '{}: duplicate undirected pairs'.format(network))
    check(qc['self_loops'] == 0, '{}: self loops'.format(network))


def components(rows):
    adjacency = defaultdict(set)
    for row in rows:
        a, b = row['taxon1_raw'], row['taxon2_raw']
        adjacency[a].add(b)
        adjacency[b].add(a)
    unseen = set(adjacency)
    result = []
    while unseen:
        start = unseen.pop()
        found, stack = {start}, [start]
        while stack:
            for neighbor in adjacency[stack.pop()]:
                if neighbor in unseen:
                    unseen.remove(neighbor)
                    found.add(neighbor)
                    stack.append(neighbor)
        result.append(found)
    return result


def is_data3(path):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook:
            preview = list(sheet.iter_rows(max_row=20, values_only=True))
            has_title = any(isinstance(v, str) and v.startswith('Supplementary Data 3.')
                            for row in preview for v in row)
            has_fields = any(set(META_FIELDS).issubset(set(row)) for row in preview)
            if has_title and has_fields:
                return True
        return False
    finally:
        workbook.close()


def inventory():
    LOG.info('STEP 1: input inventory, SHA256, and Supplementary Data 3 identification')
    paths = [('Zhou2022', RAW / 'zhou2022/NIHMS1880430-supplement-Supplementary_Table_S6.xlsx'),
             ('Guo2023', RAW / 'guo2023/1-s2.0-S193131282300029X-mmc8.xlsx')]
    paths += [('Lacruz2026', RAW / ('lacruz2026/edgeweights_{}.csv'.format(g.lower()))) for g in GROUPS]
    paths += [('Lacruz2026', RAW / 'lacruz2026/whole_cohort_data.csv')]
    expected_meta = RAW / 'lacruz2026/Lacruz2026_Supplementary_Data3.xlsx'
    meta = expected_meta
    if not expected_meta.is_file():
        candidates = [p for p in sorted(expected_meta.parent.glob('*.xlsx')) if is_data3(p)]
        STATE['metadata_candidates'] = [p.name for p in candidates]
        if len(candidates) == 1:
            meta = candidates[0]
    paths.append(('Lacruz2026', meta))
    for paper, path in paths:
        row = dict(paper=paper, file=path.name, absolute_path=str(path),
                   size_bytes=path.stat().st_size if path.is_file() else NA,
                   sha256=sha256(path) if path.is_file() else NA,
                   status=('present_alternate_filename' if path == meta and meta != expected_meta else 'present')
                   if path.is_file() else 'missing')
        STATE['inventory'].append(row)
        LOG.info('%s: %s, bytes=%s, SHA256=%s', row['status'], path.name, row['size_bytes'], row['sha256'])
    write_tsv(REPORTS / '01_input_inventory.tsv', STATE['inventory'],
              ['paper', 'file', 'absolute_path', 'size_bytes', 'sha256', 'status'])
    check(all(row['status'] != 'missing' for row in STATE['inventory']), 'Required input missing or Supplementary Data 3 ambiguous')
    check(is_data3(meta), 'Workbook content does not identify Supplementary Data 3')
    STATE['metadata_file'] = meta.name
    return paths[0][1], paths[1][1], meta


def parse_zhou(path, stage):
    """Read only the two named sheets. Apply BOTH fixed thresholds, then remove
    entire undirected components with exactly two nodes, retaining all others.
    Count assertions are evaluated before releasing either final network.
    """
    LOG.info('STEP 2: Zhou published dense microbial networks')
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    output = []
    expected_header = ('from', 'to', 'uncorrected_r', 'age_bmi_sex_corrected_r', 'pvalue', 'fdr')
    try:
        for group, (expected_edges, expected_taxa) in ZHOU_EXPECTED.items():
            name = '{} microbial network'.format(group)
            LOG.info('Zhou: reading %s', name)
            iterator = workbook[name].iter_rows(values_only=True)
            check(tuple(next(iterator)) == expected_header, '{}: unexpected columns'.format(name))
            selected, original_pairs = [], Counter()
            original_count = r_only = blanks = original_loops = 0
            for line, values in enumerate(iterator, 2):
                if all(v is None for v in values):
                    blanks += 1
                    continue
                check(len(values) == 6, '{} row {}: malformed row'.format(name, line))
                a, b, uncorrected, corrected, pvalue, qvalue = values
                a, b = taxon(a, name), taxon(b, name)
                un = num(uncorrected, name, correlation=True)
                co = num(corrected, name, correlation=True)
                num(pvalue, name, probability=True)
                q = num(qvalue, name, probability=True)
                original_count += 1
                original_pairs[tuple(sorted((a, b)))] += 1
                original_loops += a == b
                r_only += abs(co) >= 0.4
                if abs(co) >= 0.4 and q < 0.05:
                    selected.append(edge('Zhou2022', 'Zhou2022_' + group, group, path, a, b, corrected,
                                         'SparCC_covariate_corrected_correlation', 'SparCC',
                                         uncorrected_r=uncorrected, corrected_r=corrected,
                                         pvalue=pvalue, qvalue=qvalue))
            comps = components(selected)
            two_node = [c for c in comps if len(c) == 2]
            removed_nodes = set().union(*two_node) if two_node else set()
            final = [row for row in selected if row['taxon1_raw'] not in removed_nodes]
            before_qc, final_qc = edge_qc(selected), edge_qc(final)
            qc = dict(network_id='Zhou2022_' + group, source_sheet=name,
                      sheet_original_rows=original_count, sheet_blank_rows=blanks,
                      after_abs_r04=r_only, after_abs_r04_fdr005=len(selected),
                      two_node_components=len(two_node), edges_removed_two_node_components=len(selected)-len(final),
                      component_size_distribution_before_removal=json.dumps(dict(sorted(Counter(map(len, comps)).items()))),
                      final_edges=len(final), final_taxa=final_qc['unique_raw_taxa'],
                      expected_edges=expected_edges, expected_taxa=expected_taxa,
                      original_duplicate_raw_pairs=sum(n-1 for n in original_pairs.values()),
                      original_self_loops=original_loops,
                      filtered_duplicate_raw_pairs=before_qc['duplicate_raw_pairs'],
                      filtered_self_loops=before_qc['self_loops'], **final_qc)
            STATE['zhou'].append(qc)
            qc_tsv('zhou2022_parsing_qc.tsv', STATE['zhou'])
            LOG.info('Zhou %s: original=%d, |r|>=0.4=%d, both thresholds=%d, 2-node components=%d, removed edges=%d, final=%d edges/%d taxa',
                     group, original_count, r_only, len(selected), len(two_node), len(selected)-len(final), len(final), qc['final_taxa'])
            check(len(final) == expected_edges and qc['final_taxa'] == expected_taxa,
                  '{}: observed {}/{} edges/taxa; expected {}/{}'.format(group, len(final), qc['final_taxa'], expected_edges, expected_taxa))
            assert_simple(before_qc, qc['network_id'] + ' thresholded')
            assert_simple(final_qc, qc['network_id'])
            check(all(len(c) != 2 for c in components(final)), 'Zhou: a 2-node component remains')
            check(qc['original_duplicate_raw_pairs'] == 0 and original_loops == 0, 'Zhou original duplicate/self-loop anomaly')
            write_tsv(stage / ('zhou2022_{}_final_edges.tsv'.format(group)), final, FIELDS)
            output.extend(final)
    finally:
        workbook.close()
    return output


def parse_guo(path, stage):
    """Locate the explicit Edgelists row and seven Taxa1/Taxa2/Weight triples.
    Associate each triple with its exact published label in the Indexes row.
    The Jaccard and metrics sections are never treated as edge records. Retain
    all seven lists for audit; only final aggregated lists enter the denominator.
    """
    LOG.info('STEP 3: Guo seven explicitly delimited edge-list blocks')
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        grid = list(workbook['Table S7'].iter_rows(values_only=True))
    finally:
        workbook.close()
    edge_headers = [i for i, row in enumerate(grid) if row[0] == 'Edgelists']
    check(len(edge_headers) == 1, 'Guo: expected one Edgelists section')
    header = edge_headers[0]
    index_headers = [i for i, row in enumerate(grid[:header]) if row[0] == 'Indexes']
    check(len(index_headers) == 1, 'Guo: expected one Indexes network-label row')
    label_row = grid[index_headers[0]]
    starts = [c for c, value in enumerate(grid[header]) if value == 'Taxa1']
    check(len(starts) == 7, 'Guo: expected seven edge-list triples')
    blocks = {}
    for col in starts:
        check(grid[header][col:col+3] == ('Taxa1', 'Taxa2', 'Weight'), 'Guo: malformed edge header')
        check(label_row[col] not in blocks, 'Guo: duplicate network label')
        blocks[label_row[col]] = col
    check(set(blocks) == {s[1] for s in GUO_SPECS}, 'Guo: unexpected network labels')
    metrics = {row[0]: row for row in grid[:header] if row[0] in
               ['Number of nodes', 'Number of edges', 'Number of positive edges', 'Number of negative edges']}
    check(len(metrics) == 4, 'Guo: missing published QC metrics')
    all_audit, final_output, by_name = [], [], {}
    for name, published_label, suffix, condition, expected, role in GUO_SPECS:
        col = blocks[published_label]
        records, source_lines = [], []
        ended = False
        for i in range(header + 1, len(grid)):
            values = grid[i][col:col+3]
            if all(v is None for v in values):
                ended = True
                continue
            check(not ended, '{}: nonempty edge after a blank row at {}'.format(name, i+1))
            check(len(values) == 3 and all(v is not None for v in values), '{}: incomplete edge at {}'.format(name, i+1))
            a, b, weight = values
            records.append(edge('Guo2023', 'Guo2023_' + suffix, condition, path, a, b, weight,
                                'FlashWeave_association_weight', 'FlashWeave-S',
                                'primary' if role == 'final_aggregated' else 'audit_only'))
            source_lines.append(i+1)
        qc = dict(network_id='Guo2023_' + suffix, network_name=name, network_role=role,
                  source_sheet='Table S7', edge_header_row=header+1, first_column=col+1,
                  first_edge_row=source_lines[0] if source_lines else NA,
                  last_edge_row=source_lines[-1] if source_lines else NA,
                  expected_edges=expected, published_edges=metrics['Number of edges'][col],
                  published_taxa=metrics['Number of nodes'][col],
                  published_positive_edges=metrics['Number of positive edges'][col],
                  published_negative_edges=metrics['Number of negative edges'][col], **edge_qc(records))
        STATE['guo'].append(qc)
        qc_tsv('guo2023_parsing_qc.tsv', STATE['guo'])
        LOG.info('Guo %s: %d edges, %d taxa, +%d/-%d', name, len(records), qc['unique_raw_taxa'], qc['positive_edges'], qc['negative_edges'])
        check(len(records) == expected, '{}: observed {} edges, expected {}'.format(name, len(records), expected))
        for observed, published in [('raw_edges', 'published_edges'), ('unique_raw_taxa', 'published_taxa'),
                                    ('positive_edges', 'published_positive_edges'), ('negative_edges', 'published_negative_edges')]:
            check(qc[observed] == qc[published], '{}: {} does not match the workbook metric'.format(name, observed))
        assert_simple(qc, name)
        for record, line in zip(records, source_lines):
            all_audit.append(dict(record, network_name=name, network_role=role, source_sheet='Table S7', source_row=line))
        by_name[name] = records
        if role == 'final_aggregated':
            final_output.extend(records)
    # Cross-check published aggregation using exact unordered raw pairs, without
    # changing endpoint strings or substituting less-rounded intermediate weights.
    for group in ['Control', 'ME/CFS']:
        common = {pair(r) for r in by_name['Common Network']}
        unique = {pair(r) for r in by_name[group + ' Unique Edge Network']}
        aggregated = {pair(r) for r in by_name['Aggregated ' + group + ' Network']}
        check(not common.intersection(unique), 'Guo: common/validated-unique overlap')
        check(aggregated == common.union(unique), 'Guo: published aggregated edge-set mismatch')
    for name, _, suffix, _, _, role in GUO_SPECS:
        if role == 'final_aggregated':
            write_tsv(stage / ('guo2023_{}_final_edges.tsv'.format(suffix)), by_name[name], SHORT_FIELDS)
    write_tsv(stage / 'guo2023_all_networks_audit.tsv', all_audit,
              SHORT_FIELDS + ['network_name', 'network_role', 'source_sheet', 'source_row'])
    STATE['guo_aggregation_identity'] = 'PASS: aggregated = common union validated_unique; no shared pairs between those components'
    return final_output


def parse_lacruz(stage):
    """CSV files are already final SPIEC-EASI networks. Keep every record and
    retain weight text and full taxonomy strings exactly. No network estimation.
    MHO is kept in full with exploratory status because the supplied brief gives
    n=19; metadata subsequently checks this directly against Data 3.
    """
    LOG.info('STEP 4: Lacruz final published SPIEC-EASI edges, without refitting')
    all_records = []
    for group, expected in LACRUZ_EXPECTED.items():
        path = RAW / ('lacruz2026/edgeweights_{}.csv'.format(group.lower()))
        records = []
        with path.open(encoding='utf-8-sig', newline='') as handle:
            reader = csv.reader(handle)
            check(next(reader) == ['v1', 'v2', 'edgeweight'], '{}: unexpected CSV columns'.format(path.name))
            for line, values in enumerate(reader, 2):
                check(len(values) == 3 and all(v != '' for v in values), '{} row {}: incomplete edge'.format(path.name, line))
                a, b, value = values
                records.append(edge('Lacruz2026', 'Lacruz2026_' + group, group, path, a, b, value,
                                    'SPIEC_EASI_conditional_independence_weight', 'SPIEC-EASI_MB',
                                    'exploratory' if group == 'MHO' else 'primary'))
        qc = dict(network_id='Lacruz2026_' + group, condition=group, source_file=path.name,
                  primary_status='exploratory' if group == 'MHO' else 'primary',
                  expected_edges=expected, **edge_qc(records))
        STATE['lacruz'].append(qc)
        qc_tsv('lacruz2026_parsing_qc.tsv', STATE['lacruz'])
        LOG.info('Lacruz %s: %d edges, %d taxa, +%d/-%d/zero=%d, %s', group, len(records), qc['unique_raw_taxa'],
                 qc['positive_edges'], qc['negative_edges'], qc['zero_edges'], qc['primary_status'])
        check(len(records) == expected, '{}: observed {} edges, expected {}'.format(group, len(records), expected))
        assert_simple(qc, group)
        write_tsv(stage / ('lacruz2026_{}_final_edges.tsv'.format(group)), records, SHORT_FIELDS)
        all_records.extend(records)
    check(len(all_records) == 2358, 'Lacruz total must be 2358')
    return all_records


def parse_metadata(path, stage):
    """Use only explicit Data 3 phenotype/source fields. Keep the original four
    columns and Excel row provenance. Study names are not converted to projects
    or invented dataset identifiers. Manifest run_accession cells stay intact;
    the source summary additionally counts literal semicolon-delimited run IDs.
    Whole-cohort data is used ONLY for an exact header-ID inventory comparison.
    """
    LOG.info('STEP 5: Lacruz Supplementary Data 3 source metadata; no exclusion')
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    manifest = []
    try:
        for sheet in workbook:
            grid = list(sheet.iter_rows(values_only=True))
            candidates = [i for i, row in enumerate(grid) if set(META_FIELDS).issubset(set(row))]
            check(len(candidates) == 1, '{}: cannot uniquely identify metadata columns'.format(sheet.title))
            header_row = candidates[0]
            columns = list(grid[header_row])
            check(all(isinstance(c, str) and c for c in columns) and len(set(columns)) == len(columns),
                  '{}: missing or repeated column name'.format(sheet.title))
            STATE['sheets'].append(dict(source_file=path.name, sheet_name=sheet.title,
                                       header_row=header_row+1, column_names=json.dumps(columns, ensure_ascii=False)))
            LOG.info('Data 3 sheet=%s header_row=%d columns=%s', sheet.title, header_row+1, columns)
            for i, values in enumerate(grid[header_row+1:], header_row+2):
                if all(v is None for v in values):
                    continue
                record = dict(zip(columns, values))
                check(all(record[f] is not None and str(record[f]) != '' for f in META_FIELDS),
                      'Data 3 row {}: missing key field'.format(i))
                check(record['phenotype'] in GROUPS, 'Data 3 row {}: unknown explicit phenotype {!r}'.format(i, record['phenotype']))
                check(isinstance(record['run_accession'], str), 'Data 3: non-string accession')
                record.update(paper_id='Lacruz2026', network_id='Lacruz2026_' + record['phenotype'],
                              source_file=path.name, source_sheet=sheet.title, source_row=i,
                              primary_status='exploratory' if record['phenotype'] == 'MHO' else 'primary')
                manifest.append(record)
    finally:
        workbook.close()
    check(manifest, 'Data 3: no sample records')
    subject_ids = [str(r['subject_id']) for r in manifest]
    check(len(set(subject_ids)) == len(subject_ids), 'Data 3: duplicate subject IDs')
    phenotype_counts = Counter(r['phenotype'] for r in manifest)
    study_counts = Counter(r['study_name'] for r in manifest)
    distribution = Counter((r['study_name'], r['phenotype']) for r in manifest)
    check(phenotype_counts['MHO'] == 19, 'Data 3: MHO n differs from 19; no phenotype correction attempted')
    original_columns = list(dict.fromkeys(c for row in STATE['sheets'] for c in json.loads(row['column_names'])))
    write_tsv(stage / 'lacruz2026_sample_source_manifest.tsv', manifest,
              original_columns + ['paper_id', 'network_id', 'source_file', 'source_sheet', 'source_row', 'primary_status'])
    qc_tsv('lacruz2026_supplementary_data3_schema.tsv', STATE['sheets'])
    # Two source levels in one table; they are separate aggregations and MUST NOT
    # be summed together. Per-run counts can also exceed sample totals, since a
    # sample may have multiple runs. No run is converted to a BioProject.
    study_subjects, run_subjects = defaultdict(set), defaultdict(set)
    for row in manifest:
        group, study, subject = row['phenotype'], row['study_name'], str(row['subject_id'])
        study_subjects[(group, study)].add(subject)
        tokens = row['run_accession'].split(';')
        check(all(tokens), 'Data 3: empty run-accession token')
        check(len(tokens) == len(set(tokens)), 'Data 3: repeated run accession within one sample')
        for accession in tokens:
            # No strip/case conversion and no prefix-based project inference.
            run_subjects[(group, study, accession)].add(subject)
    summary = []
    for (group, study), ids in sorted(study_subjects.items()):
        summary.append(dict(network_id='Lacruz2026_' + group, condition=group,
                            primary_status='exploratory' if group == 'MHO' else 'primary',
                            source_level='study', source_field='study_name', study=study,
                            dataset=NA, accession=NA, accession_type=NA, bioproject=NA, sample_count=len(ids)))
    for (group, study, accession), ids in sorted(run_subjects.items()):
        summary.append(dict(network_id='Lacruz2026_' + group, condition=group,
                            primary_status='exploratory' if group == 'MHO' else 'primary',
                            source_level='run_accession', source_field='run_accession', study=study,
                            dataset=NA, accession=accession, accession_type='run', bioproject=NA, sample_count=len(ids)))
    write_tsv(stage / 'lacruz2026_network_source_summary.tsv', summary,
              ['network_id', 'condition', 'primary_status', 'source_level', 'source_field', 'study',
               'dataset', 'accession', 'accession_type', 'bioproject', 'sample_count'])
    for group in GROUPS:
        check(sum(r['sample_count'] for r in summary if r['source_level'] == 'study' and r['condition'] == group)
              == phenotype_counts[group], 'Source summary sample-count mismatch')
    cohort_path = RAW / 'lacruz2026/whole_cohort_data.csv'
    with cohort_path.open(encoding='utf-8-sig', newline='') as handle:
        header = next(csv.reader(handle))  # Never read abundance rows to build a graph.
    check(header[0] == 'taxa', 'whole_cohort_data.csv: unexpected leading header')
    cohort_ids = header[1:]
    check(len(set(cohort_ids)) == len(cohort_ids), 'whole_cohort_data.csv: duplicate header sample IDs')
    source_set, cohort_set = set(subject_ids), set(cohort_ids)
    differences = []
    for record in manifest:
        if str(record['subject_id']) not in cohort_set:
            differences.append(dict(origin='Supplementary_Data3_only', id_raw=record['subject_id'],
                                    study_name=record['study_name'], phenotype=record['phenotype'],
                                    source_file=path.name, source_row=record['source_row']))
    for subject in sorted(cohort_set - source_set):
        differences.append(dict(origin='whole_cohort_header_only', id_raw=subject, study_name=NA,
                                phenotype=NA, source_file=cohort_path.name, source_row=1))
    write_tsv(REPORTS / 'lacruz2026_sample_id_exact_match_qc.tsv', differences,
              ['origin', 'id_raw', 'study_name', 'phenotype', 'source_file', 'source_row'])
    if differences:
        warning = ('Lacruz sample IDs: {} exact matches; {} Data 3 IDs and {} whole-cohort header IDs do not match. '
                   'Unmatched Data 3 studies: {}. IDs are preserved without inferred reconciliation.').format(
                       len(source_set & cohort_set), len(source_set-cohort_set), len(cohort_set-source_set),
                       dict(Counter(r['study_name'] for r in differences if r['origin'] == 'Supplementary_Data3_only')))
        STATE['warnings'].append(warning)
        LOG.warning(warning)
    STATE['metadata'] = dict(sample_records=len(manifest), unique_subjects=len(source_set),
                             phenotype_counts=dict(phenotype_counts), study_counts=dict(study_counts),
                             study_phenotype=[dict(study=study, **{g: distribution[(study, g)] for g in GROUPS},
                                                   total=study_counts[study]) for study in sorted(study_counts)],
                             missing_key_fields=0, duplicate_subject_ids=0,
                             explicit_dataset_field=False, explicit_bioproject_field=False,
                             cohort_sample_columns=len(cohort_ids), exact_matches=len(source_set & cohort_set),
                             manifest_only=len(source_set-cohort_set), cohort_only=len(cohort_set-source_set),
                             run_source_rows=len(run_subjects), unique_run_accessions=len({key[2] for key in run_subjects}))
    LOG.info('Data 3: %d samples; phenotypes=%s; studies=%s', len(manifest), dict(phenotype_counts), dict(study_counts))


def table(rows, columns):
    if not rows:
        return '_尚未完成。_\n'
    def escape(value):
        return str(value).replace('|', '\\|').replace('\n', ' ')
    return '\n'.join(['| ' + ' | '.join(columns) + ' |', '| ' + ' | '.join(['---'] * len(columns)) + ' |'] +
                     ['| ' + ' | '.join(escape(row.get(c, NA)) for c in columns) + ' |' for row in rows]) + '\n'


def report():
    lines = ['# 第一阶段：published external network parsing QC', '',
             '状态：**{}**；UTC：{}。'.format(STATE['status'], datetime.now(timezone.utc).isoformat()), '',
             '范围仅为本地文件核查、作者最终 edge lists 解析、source metadata 和 QC。未联网或下载文件；'
             '未修改 raw；未执行 taxonomy/GTDB/NCBI mapping、MACK-DB overlap、source exclusion 或重新构网。', '',
             '复现：`python3 -s scripts/01_parse_external_networks.py`（在固定 BASE 下运行）。'
             'Python optimized 模式被禁止，以保证 assert 生效。系统 openpyxl={}。'.format(openpyxl.__version__), '',
             '## 输入清单与 SHA256', '',
             table(STATE['inventory'], ['paper', 'file', 'size_bytes', 'sha256', 'status']),
             '实际 Supplementary Data 3：`{}`；通过 workbook 标题及字段确认。'.format(STATE.get('metadata_file', NA)), '',
             '## Zhou2022', '',
             '仅读取 `MS microbial network` 与 `HHC microbial network`。先同时筛选 '
             '`abs(age_bmi_sex_corrected_r) >= 0.4` 和 `fdr < 0.05`，再建立无向图，'
             '仅删除恰好 2 个 nodes 的整个 connected components。≥3 nodes components 全部保留；未调整阈值。', '',
             table(STATE['zhou'], ['network_id', 'sheet_original_rows', 'after_abs_r04', 'after_abs_r04_fdr005',
                                  'two_node_components', 'edges_removed_two_node_components', 'final_edges', 'final_taxa',
                                  'positive_edges', 'negative_edges', 'zero_edges', 'duplicate_raw_pairs', 'self_loops']),
             '`after_abs_r04` 是仅按 |r| 筛选的记录数；`after_abs_r04_fdr005` 是两个阈值同时应用后的记录数。'
             '严格预期 MS=6742 edges/786 taxa，HHC=5688 edges/773 taxa。', '',
             '## Guo2023', '',
             '按照 `Edgelists` 行的七组 `Taxa1 / Taxa2 / Weight` 及同列 `Indexes` 网络名称精确定位。'
             'Jaccard 区域不参与解析；Network metrics 仅用于对节点数、边数和正负边数交叉核对。'
             '所有 7 个网络保存在 audit 文件，主 denominator 只使用最终两个 Aggregated 网络。', '',
             table(STATE['guo'], ['network_name', 'network_role', 'raw_edges', 'unique_raw_taxa', 'positive_edges',
                                 'negative_edges', 'zero_edges', 'duplicate_raw_pairs', 'self_loops']),
             STATE.get('guo_aggregation_identity', 'Aggregation identity: not yet checked.'), '',
             '最终 AggControl=388，AggMECFS=358；未使用 210/223 作为主 denominator。'
             '保留最终列表自身的 Weight 精度，不用 intermediate weight 替换。'
             'pvalue/qvalue 均为 NA：最终 edge lists 未逐 edge 提供，不生成或猜测。'
             'audit 中其余 5 个网络 primary_status=audit_only。', '',
             '## Lacruz2026', '',
             '直接保留作者四个最终 CSV 中的全部 edge records。保留 edgeweight 原始数值文本和完整 taxonomy 字符串。'
             '未拆 taxonomy、删除 unnamed SGB 或重新运行 SPIEC-EASI。pvalue/qvalue 均为 NA。', '',
             table(STATE['lacruz'], ['network_id', 'primary_status', 'raw_edges', 'unique_raw_taxa', 'positive_edges',
                                    'negative_edges', 'zero_edges', 'duplicate_raw_pairs', 'self_loops']),
             'MHO 完整保留且标记 exploratory；其 n=19 由本地 Supplementary Data 3 直接核对。', '',
             '## Lacruz Supplementary Data 3', '',
             table(STATE['sheets'], ['source_file', 'sheet_name', 'header_row', 'column_names'])]
    metadata = STATE.get('metadata')
    if metadata:
        lines += ['识别：subject/sample ID=`subject_id`；sequencing run accession=`run_accession`；'
                  'study=`study_name`；phenotype classification=`phenotype`（MHNO/MHO/MUNO/MUO 原值）。'
                  '没有独立 dataset、BioProject/project 字段；summary 相应字段为 NA，未推断。', '',
                  table(metadata['study_phenotype'], ['study'] + GROUPS + ['total']),
                  '合计 {} 条、{} 个独立 subject；phenotype 分布：{}。四个关键字段均无缺失，subject ID 无重复。'.format(
                      metadata['sample_records'], metadata['unique_subjects'], json.dumps(metadata['phenotype_counts'])), '',
                  'manifest 保留原四列及 source_file/source_sheet/source_row 等来源字段。数值型 Excel subject ID '
                  '按其存储数值输出；不添加前后缀。原始分号连接 run_accession 完整保存在 manifest。', '',
                  'network_source_summary 含 `source_level=study` 与 `run_accession` 两种统计。'
                  '仅为 run 层统计按原有分号拆分 accession；不推断 BioProject。sample_count 为对应来源的独立 subject 数。'
                  '两种层级不可相加；一个 sample 可有多个 runs，所以 run 层计数不可当作总 sample 数。'
                  '独立 run accession 数={}。'.format(metadata['unique_run_accessions']), '',
                  'whole_cohort_data.csv 仅读取表头核查 sample IDs；{} 个 sample columns，精确匹配 {} 个；'
                  'Data 3 独有 {} 个，cohort 表头独有 {} 个。未将未匹配 ID 人工或模糊配对，未改变网络或 phenotype。'
                  '差异见 `lacruz2026_sample_id_exact_match_qc.tsv`。'.format(
                      metadata['cohort_sample_columns'], metadata['exact_matches'], metadata['manifest_only'], metadata['cohort_only']), '']
    lines += ['## 合并与严格 QC', '',
              '合并列顺序：`' + ' / '.join(FIELDS) + '`。缺失字段为 NA；taxon 字符串和端点顺序不改。', '',
              'duplicate_raw_pairs 在**每个网络内部**按无向原始字符串 pair 计算，数值为多余记录数 '
              '(记录数−独立 pair 数)。没有跨网络去重；报告 taxa 也是原始 label 数，尚未作生物学名称归并。'
              '检查了空值、数值有限性、符号一致性、Zhou correlation 和 p/q 范围、重复 pair 和 self-loop。', '']
    if 'merged' in STATE:
        lines += [table([STATE['merged']], ['Zhou2022', 'Guo2023', 'Lacruz2026', 'all_edge_records', 'primary_edge_records', 'exploratory_edge_records']),
                  '预期：12430 + 746 + 2358 = **15534**；primary=**14836**；exploratory=**698**。'
                  '以上为 network-level records，不跨 network 去重。', '',
                  '所有待发布 TSV 已回读并与内存记录逐字段核对。raw 输入结束时重新计算 SHA256 并与开始值逐一比较。', '']
    lines += ['## 异常与限制', '']
    lines += ['- ' + warning for warning in STATE['warnings']] or ['- 未发现输入或解析异常。']
    lines += ['- 默认用户 Python 环境的 openpyxl/NumPy 不兼容；使用已有系统环境 `python3 -s` 完成解析，未安装软件。']
    if 'error' in STATE:
        lines += ['', '**停止原因：** ' + STATE['error'], '',
                  '本次 staging 中的 parsed 文件未发布。若目标目录有先前运行产物，不应视为本次有效结果。']
    lines += ['', '## 输出文件', ''] + ['- `' + name + '`' for name in STATE['outputs']]
    lines += ['', '第一阶段完成后停止；没有开始第二阶段。' if STATE['status'].startswith('PASS') else '未进入第二阶段。', '']
    (REPORTS / '01_parsing_qc.md').write_text('\n'.join(lines), encoding='utf-8')


def verify_raw_unchanged():
    LOG.info('Rechecking input size and SHA256 against the initial inventory')
    checks = []
    for row in STATE['inventory']:
        if row['status'] == 'missing':
            continue
        path = Path(row['absolute_path'])
        after = sha256(path) if path.is_file() else 'MISSING'
        size = path.stat().st_size if path.is_file() else NA
        checks.append(dict(file=row['file'], sha256_before=row['sha256'], sha256_after=after,
                           size_bytes_before=row['size_bytes'], size_bytes_after=size,
                           unchanged=after == row['sha256'] and size == row['size_bytes']))
    if checks:
        qc_tsv('01_input_integrity_qc.tsv', checks)
        check(all(row['unchanged'] for row in checks), 'Raw input content changed during execution')
    STATE['raw_unchanged'] = True


def main():
    if sys.flags.optimize:
        raise RuntimeError('Do not use python -O/-OO: required assertions must remain enabled')
    for directory in [PARSED, REPORTS]:
        directory.mkdir(parents=True, exist_ok=True)
    LOG.setLevel(logging.INFO)
    formatter = logging.Formatter('%(asctime)s %(levelname)s %(message)s')
    for handler in [logging.FileHandler(REPORTS / '01_parse_external_networks.log', mode='w', encoding='utf-8'),
                    logging.StreamHandler(sys.stdout)]:
        handler.setFormatter(formatter)
        LOG.addHandler(handler)
    LOG.info('START phase 1; BASE=%s; Python=%s; openpyxl=%s; raw is read-only', BASE, sys.version.split()[0], openpyxl.__version__)
    try:
        zhou_path, guo_path, metadata_path = inventory()
        # Do not publish a partial merged result if any edge-count assertion fails.
        with tempfile.TemporaryDirectory(prefix='.phase1-', dir=str(PARSED)) as staging:
            stage = Path(staging)
            combined = parse_zhou(zhou_path, stage)
            combined.extend(parse_guo(guo_path, stage))
            combined.extend(parse_lacruz(stage))
            parse_metadata(metadata_path, stage)
            LOG.info('STEP 6: unified final/exploratory edges and denominator assertions')
            papers = Counter(row['paper_id'] for row in combined)
            statuses = Counter(row['primary_status'] for row in combined)
            networks = Counter(row['network_id'] for row in combined)
            STATE['merged'] = dict(papers, all_edge_records=len(combined),
                                   primary_edge_records=statuses['primary'], exploratory_edge_records=statuses['exploratory'])
            check(papers == {'Zhou2022': 12430, 'Guo2023': 746, 'Lacruz2026': 2358}, 'Per-paper denominator mismatch')
            check(len(combined) == 15534, 'Merged edge records must be exactly 15534')
            check(statuses == {'primary': 14836, 'exploratory': 698}, 'Primary/exploratory denominator mismatch')
            expected_networks = {'Zhou2022_MS': 6742, 'Zhou2022_HHC': 5688,
                                 'Guo2023_AggControl': 388, 'Guo2023_AggMECFS': 358}
            expected_networks.update({'Lacruz2026_' + group: n for group, n in LACRUZ_EXPECTED.items()})
            check(networks == expected_networks, 'Merged network-level counts mismatch')
            for row in combined:
                check(row['sign'] == sign(num(row['association_value'], row['network_id'])), 'Sign inconsistency')
                if row['paper_id'] != 'Zhou2022':
                    check(all(row[f] == NA for f in ['uncorrected_r', 'corrected_r', 'pvalue', 'qvalue']), 'Invented missing numeric field')
                if row['primary_status'] == 'exploratory':
                    check(row['network_id'] == 'Lacruz2026_MHO', 'Unexpected exploratory network')
            write_tsv(stage / 'all_new_external_edges_raw.tsv', combined, FIELDS)
            LOG.info('Unified records=%d; primary=%d; exploratory=%d; all staged TSV read-back checks passed',
                     len(combined), statuses['primary'], statuses['exploratory'])
            verify_raw_unchanged()
            LOG.info('STEP 7: publish validated parsed TSVs and final report')
            for source in sorted(stage.iterdir()):
                target = PARSED / source.name
                os.replace(str(source), str(target))
                STATE['outputs'].append(str(target.relative_to(BASE)))
        STATE['status'] = 'PASS_WITH_WARNING' if STATE['warnings'] else 'PASS'
        STATE['outputs'] += ['reports/' + p.name for p in sorted(REPORTS.iterdir()) if p.is_file()]
        if 'reports/01_parsing_qc.md' not in STATE['outputs']:
            STATE['outputs'].append('reports/01_parsing_qc.md')
        STATE['outputs'].append('scripts/01_parse_external_networks.py')
        report()
        LOG.info('DONE %s: exactly 15534 records (14836 primary + 698 exploratory); phase 2 not started', STATE['status'])
        return 0
    except Exception as error:
        STATE['status'] = 'FAIL'
        STATE['error'] = '{}: {}'.format(type(error).__name__, error)
        LOG.exception('STOP: %s', STATE['error'])
        if not STATE.get('raw_unchanged'):
            try:
                verify_raw_unchanged()
            except Exception as integrity_error:
                STATE['warnings'].append('Raw integrity check failed: {}'.format(integrity_error))
                LOG.exception('Raw integrity check failed')
        report()
        return 1


if __name__ == '__main__':
    sys.exit(main())
